"""Document ingestion service: parse → chunk → embed → store in Qdrant."""
import os
import uuid
import json
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import fitz  # PyMuPDF
from docx import Document as DocxDocument
import markdown as md_lib

from langchain_text_splitters import RecursiveCharacterTextSplitter

from database.qdrant_client import get_qdrant_client, init_collection

UPLOADS_DIR = Path(__file__).parent.parent.parent / "uploads"
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

METADATA_FILE = UPLOADS_DIR / "metadata.json"

CHUNK_SIZE    = 600
CHUNK_OVERLAP = 150

splitter = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
)


# ─── Metadata persistence ─────────────────────────────────────────────────────

def _load_metadata() -> dict:
    if METADATA_FILE.exists():
        return json.loads(METADATA_FILE.read_text())
    return {}


def _save_metadata(data: dict) -> None:
    METADATA_FILE.write_text(json.dumps(data, indent=2))


def list_documents() -> list[dict]:
    meta = _load_metadata()
    return list(meta.values())


def get_document(doc_id: str) -> Optional[dict]:
    return _load_metadata().get(doc_id)


# ─── Text extraction ──────────────────────────────────────────────────────────

def _extract_text_pdf(path: Path) -> list[tuple[str, int]]:
    """Returns list of (text, page_number)."""
    pages = []
    doc = fitz.open(str(path))
    for i, page in enumerate(doc, 1):
        text = page.get_text().strip()
        if text:
            pages.append((text, i))
    return pages


def _extract_text_docx(path: Path) -> list[tuple[str, int]]:
    doc = DocxDocument(str(path))
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [(full_text, 1)]


def _extract_text_txt(path: Path) -> list[tuple[str, int]]:
    return [(path.read_text(encoding="utf-8", errors="ignore"), 1)]


def _extract_text_md(path: Path) -> list[tuple[str, int]]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    # Strip markdown to plain text (simple approach)
    import re
    plain = re.sub(r"[#*`_~\[\]()]", "", raw)
    return [(plain, 1)]


def _extract_pages(path: Path, suffix: str) -> list[tuple[str, int]]:
    s = suffix.lower()
    if s == ".pdf":
        return _extract_text_pdf(path)
    elif s in (".docx",):
        return _extract_text_docx(path)
    elif s in (".md", ".markdown"):
        return _extract_text_md(path)
    else:
        return _extract_text_txt(path)


# ─── Ingest ───────────────────────────────────────────────────────────────────

async def ingest_document(file_bytes: bytes, filename: str) -> dict:
    """Parse, chunk, embed, store. Returns metadata dict."""
    doc_id  = str(uuid.uuid4())
    suffix  = Path(filename).suffix
    dest    = UPLOADS_DIR / f"{doc_id}{suffix}"
    dest.write_bytes(file_bytes)

    # Extract text per page
    pages = _extract_pages(dest, suffix)

    # Build chunks with page metadata
    all_chunks: list[tuple[str, int]] = []   # (text, page_number)
    for page_text, page_num in pages:
        splits = splitter.split_text(page_text)
        for s in splits:
            all_chunks.append((s, page_num))

    if not all_chunks:
        raise ValueError("No text could be extracted from the document.")

    texts       = [c[0] for c in all_chunks]
    page_nums   = [c[1] for c in all_chunks]

    # Store in Qdrant
    client = get_qdrant_client()
    init_collection("documents")

    from qdrant_client.models import PointStruct, Document

    points = []
    for i, text in enumerate(texts):
        # Generate a valid UUID from doc_id and index deterministically
        chunk_id = str(uuid.uuid5(uuid.UUID(doc_id), str(i)))
        points.append(
            PointStruct(
                id=chunk_id,
                vector=Document(
                    text=text,
                    model="sentence-transformers/all-MiniLM-L6-v2",
                ),
                payload={
                    "document_id": doc_id,
                    "filename":    filename,
                    "page_number": page_nums[i],
                    "chunk_index": i,
                    "content":     text,
                },
            )
        )

    client.upsert(
        collection_name="documents",
        points=points,
    )

    # Persist document metadata
    doc_meta = {
        "id":           doc_id,
        "filename":     filename,
        "file_type":    suffix.lstrip(".").upper(),
        "chunk_count":  len(texts),
        "uploaded_at":  datetime.now(timezone.utc).isoformat(),
    }
    all_meta = _load_metadata()
    all_meta[doc_id] = doc_meta
    _save_metadata(all_meta)

    return doc_meta


async def delete_document(doc_id: str) -> bool:
    """Remove document chunks from Qdrant and metadata."""
    all_meta = _load_metadata()
    if doc_id not in all_meta:
        return False

    doc_info = all_meta[doc_id]
    chunk_count = doc_info.get("chunk_count", 0)

    # Delete chunks from Qdrant
    client = get_qdrant_client()
    from qdrant_client.models import Filter, FieldCondition, MatchValue
    try:
        client.delete(
            collection_name="documents",
            points_selector=Filter(
                must=[
                    FieldCondition(
                        key="document_id",
                        match=MatchValue(value=doc_id),
                    )
                ]
            ),
        )
    except Exception:
        pass

    # Remove file
    suffix = Path(doc_info["filename"]).suffix
    file_path = UPLOADS_DIR / f"{doc_id}{suffix}"
    if file_path.exists():
        file_path.unlink()

    del all_meta[doc_id]
    _save_metadata(all_meta)
    return True
